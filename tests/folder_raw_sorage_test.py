import pytest
import os
from docx import Document
from fpdf import FPDF
from mirage import FolderRawStorage

# Фикстура для создания тестовых файлов
@pytest.fixture(scope="module")
def test_files_dir(tmpdir_factory):
    # Создаём временную папку для тестовых файлов
    test_files_dir = tmpdir_factory.mktemp("test_files")

    # Создаём test.txt
    txt_file_path = os.path.join(test_files_dir, "test.txt")
    with open(txt_file_path, "w", encoding="utf-8") as file:
        file.write("This is a test text file.")

    # Создаём test.docx
    docx_file_path = os.path.join(test_files_dir, "test.docx")
    doc = Document()
    doc.add_paragraph("This is a test Word document.")
    doc.save(docx_file_path)

    # Создаём test.pdf
    pdf_file_path = os.path.join(test_files_dir, "test.pdf")
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, txt="This is a test PDF document.", ln=True, align="C")
    pdf.output(pdf_file_path)

    # Возвращаем путь к папке с тестовыми файлами
    return test_files_dir

# Тесты
def test_folder_raw_storage_initialization(test_files_dir):
    storage = FolderRawStorage(str(test_files_dir))
    assert len(storage.get_indexes()) == 3  # Ожидаем 3 файла: test.txt, test.pdf, test.docx

def test_folder_raw_storage_read_txt(test_files_dir):
    storage = FolderRawStorage(str(test_files_dir))
    assert storage["test.txt"] == "This is a test text file."

def test_folder_raw_storage_read_docx(test_files_dir):
    storage = FolderRawStorage(str(test_files_dir))
    assert storage["test.docx"] == "This is a test Word document."

def test_folder_raw_storage_read_pdf(test_files_dir):
    storage = FolderRawStorage(str(test_files_dir))
    assert storage["test.pdf"] == "This is a test PDF document."